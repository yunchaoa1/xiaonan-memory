# -*- coding: utf-8 -*-
"""会议纪要 Markdown → Word（python-docx；标题 / 表格 / 加粗 / 引用 / 列表）

用法：改 MD 路径（或命令行传参）→ python md2docx_meeting.py [in.md] [out.docx]
支持：# 标题 · ## 一级 · ### 二级 · | 表格 | · - 列表 · 1. 编号 · > 引用 · **加粗**内联
"""
import os, re, sys

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

MD = sys.argv[1] if len(sys.argv) > 1 else r"D:\Documents\我的文档\会议记录\2026-09-30_技术部会议记录.md"
OUT = sys.argv[2] if len(sys.argv) > 2 else os.path.splitext(MD)[0] + ".docx"

doc = Document()
st = doc.styles["Normal"]
st.font.name = "微软雅黑"
st.font.size = Pt(10.5)
for sec in doc.sections:                      # 窄边距，多塞点内容
    sec.left_margin = sec.right_margin = Pt(50)


def add_runs(p, text):
    """**加粗** 内联解析"""
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if not part:
            continue
        r = p.add_run(part)
        r.bold = (i % 2 == 1)


lines = open(MD, encoding="utf-8").read().splitlines()
i = 0
while i < len(lines):
    ln = lines[i].rstrip()
    if not ln.strip():
        i += 1
        continue

    if ln.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:\-|]+\|$", lines[i + 1].strip()):
        header = [c.strip() for c in ln.strip("|").split("|")]
        i += 2
        body = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            body.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
            i += 1
        t = doc.add_table(rows=1, cols=len(header))
        try:
            t.style = "Light Grid Accent 1"
        except Exception:
            pass
        for c, h in enumerate(header):
            cell = t.rows[0].cells[c]
            cell.text = ""
            add_runs(cell.paragraphs[0], h)
            for r in cell.paragraphs[0].runs:
                r.bold = True
        for row in body:
            cells = t.add_row().cells
            for c, v in enumerate(row[: len(header)]):
                cells[c].text = ""
                add_runs(cells[c].paragraphs[0], v)
        doc.add_paragraph()
        continue

    if ln.startswith("### "):
        doc.add_heading(ln[4:], level=3)
    elif ln.startswith("## "):
        doc.add_heading(ln[3:], level=2)
    elif ln.startswith("# "):
        h = doc.add_heading(ln[2:], level=1)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif ln.startswith("> "):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(18)
        r = p.add_run(ln[2:])
        r.italic = True
        r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    elif re.match(r"^\s*[-*] ", ln):
        p = doc.add_paragraph(style="List Bullet")
        add_runs(p, re.sub(r"^\s*[-*] ", "", ln))
    elif re.match(r"^\s*\d+\. ", ln):
        p = doc.add_paragraph(style="List Number")
        add_runs(p, re.sub(r"^\s*\d+\. ", "", ln))
    else:
        p = doc.add_paragraph()
        add_runs(p, ln)
    i += 1

doc.save(OUT)
print("已生成:", OUT)
