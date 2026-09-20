# -*- coding: utf-8 -*-
"""项目拓扑图 → WPS 演示(.pptx) 最小可跑骨架

用途：换机器/新环境时起手用；正式版在 D:\\Hermes\\scripts\\gen_topology_pptx.py
（那边是 DATA + box()/arrow() 的完整结构，改数据重跑即出图）。

运行：python topology_slide_skeleton.py 输出.pptx
要点：全部原生形状（可在 WPS 演示里双击编辑）；跑完检查"越界数"必须为 0。

凡哥的硬规则（写死在结构里，别省）：
  · 模块卡：状态 / 开始时间 / 预计交付时间
  · 超期：⚠卡点 / 解决方案 / 再次交付时间
  · 图最下方固定「待办事项」区
"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# (填充, 边框, 字色) —— 四色状态 + 卡片底
DONE  = (RGBColor(0xD9, 0xEA, 0xD3), RGBColor(0x6A, 0xA8, 0x4F), RGBColor(0x27, 0x4E, 0x13))
PROG  = (RGBColor(0xDC, 0xE9, 0xF7), RGBColor(0x44, 0x72, 0xC4), RGBColor(0x1F, 0x38, 0x64))
BLOCK = (RGBColor(0xFC, 0xE5, 0xD5), RGBColor(0xED, 0x7D, 0x31), RGBColor(0x84, 0x3C, 0x0C))
NONE  = (RGBColor(0xF2, 0xF2, 0xF2), RGBColor(0xBF, 0xBF, 0xBF), RGBColor(0x59, 0x59, 0x59))
CARD  = (RGBColor(0xFB, 0xFC, 0xFE), RGBColor(0xB8, 0xCC, 0xE8), RGBColor(0x1F, 0x38, 0x64))
TODO  = (RGBColor(0xFF, 0xF7, 0xE6), RGBColor(0xE6, 0xB8, 0x4C), RGBColor(0x7F, 0x60, 0x00))
PLAIN = (RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0x40, 0x40, 0x40))

W, H = 13.333, 7.5  # 16:9，单位英寸


def new_slide():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    return prs, prs.slides.add_slide(prs.slide_layouts[6])  # 6 = 空白版式


def box(sl, x, y, w, h, text, style=PROG, size=8, bold=False,
        align=PP_ALIGN.CENTER, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08):
    """加一个原生形状并写多行文字（\\n 分行）。"""
    fill, line, font = style
    s = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            s.adjustments[0] = radius
        except Exception:
            pass
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = line
    s.line.width = Pt(1)
    s.shadow.inherit = False
    tf = s.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, line_text in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line_text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = font
    return s


def arrow(sl, x, y, w=0.22, h=0.22, style=PROG):
    """节点之间的箭头（用块箭头形状，简单可靠）。"""
    return box(sl, x, y, w, h, "", style, shape=MSO_SHAPE.RIGHT_ARROW)


def check_bounds(sl, w=W, h=H):
    """越界检查：返回越界形状列表（交付前必须为空）。"""
    bad = []
    for s in sl.shapes:
        x, y = s.left / 914400.0, s.top / 914400.0
        ww, hh = s.width / 914400.0, s.height / 914400.0
        if x < 0 or y < 0 or x + ww > w + 0.01 or y + hh > h + 0.01:
            bad.append((round(x, 2), round(y, 2), round(x + ww, 2), round(y + hh, 2)))
    return bad


def main(out):
    prs, sl = new_slide()

    # ── 标题 + 四色图例 ─────────────────────────────
    box(sl, 4.35, 0.08, 4.6, 0.40, "技术项目主管 · 凡哥",
        (RGBColor(0x1F, 0x38, 0x64), RGBColor(0x1F, 0x38, 0x64), RGBColor(0xFF, 0xFF, 0xFF)),
        size=13, bold=True)
    for i, (t, st) in enumerate([("✅ 已完成", DONE), ("▶ 进行中", PROG),
                                 ("⚠ 卡点", BLOCK), ("○ 未开始", NONE)]):
        box(sl, 4.50 + i * 0.95, 0.54, 0.90, 0.22, t, st, size=6.5)

    # ── 项目卡（一个项目一张卡）─────────────────────
    box(sl, 0.14, 0.80, 6.34, 2.02, "", CARD)
    box(sl, 0.18, 0.84, 6.26, 0.34, "① 项目名 · ▶ 进行中", PROG, size=9, bold=True)
    # 模块卡：状态 / 开始 / 预计交付（+实际交付）
    box(sl, 0.18, 1.20, 6.26, 0.30,
        "模块A · 小X　｜　开始 2026-08　·　预计交付 2026-09-22", DONE, size=7)
    # 超期模块 → 卡点三件套
    box(sl, 0.30, 1.58, 6.02, 0.70,
        "⚠ 卡点：……\n解决方案：……　·　再次交付：待定", BLOCK, size=6.5, align=PP_ALIGN.LEFT)
    # 节点链路示例：写作 → 剧本 → 分镜脚本 …（用 arrow() 连）
    nodes = ["写作", "剧本", "分镜脚本"]
    for i, n in enumerate(nodes):
        box(sl, 0.30 + i * 1.24, 2.34, 1.02, 0.32, n, PROG, size=7)
        if i < len(nodes) - 1:
            arrow(sl, 1.34 + i * 1.24, 2.39)

    # ── 图最下方：待办事项（固定区）──────────────────
    box(sl, 0.14, 5.62, 13.02, 1.72, "", TODO)
    box(sl, 0.24, 5.66, 4.00, 0.30, "📌 待办事项", TODO, size=9, bold=True, align=PP_ALIGN.LEFT)
    for i, t in enumerate(["1. ……（谁做）", "2. ……", "3. ……"]):
        box(sl, 0.28, 5.98 + i * 0.32, 12.70, 0.30, t, PLAIN, size=6.5,
            align=PP_ALIGN.LEFT, shape=MSO_SHAPE.RECTANGLE)

    bad = check_bounds(sl)
    prs.save(out)
    print(f"已生成: {out} | 形状数: {len(sl.shapes)} | 越界数: {len(bad)}")
    for b in bad:
        print("  越界:", b)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "拓扑图.pptx")
