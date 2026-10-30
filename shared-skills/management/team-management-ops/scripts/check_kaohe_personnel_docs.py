# -*- coding: utf-8 -*-
"""校验「一人一份」的月度考核表 docx：无串人、无残留 Markdown 标记、表格齐全。

用法：
    python check_kaohe_personnel_docs.py "D:\\Documents\\我的文档\\技术部管理\\考核下发_2026-10"

判据（任一不过即为不合格，必须回炉重生成）：
  1. 残留 '**' 或 '`' 必须为 0（md2docx 正常转换不会留下这些）
  2. 每份表格数 >= 4（固定区／变动区／临时区／自评合计）
  3. 本人姓名必须在正文里
  4. 别人的「独有分区标题」不许出现在这一份里（串人）
"""
import os
import sys

# 每人的“独有分区标题”关键词 —— 换年度/换人时只改这里
EXCLUSIVE = {
    "李林":   ["影剧工坊 2.0 落地", "1.0 稳定、部署与后期完善", "服务器统计与费用监控"],
    "吴伟俊": ["浪浪椰前端", "智提分前端", "影剧工坊 2.0 前端"],
    "何锦波": ["浪浪椰后端交付", "商业化后端", "工程化与部署运维"],
    "唐光辉": ["工资条智能体", "考勤智能体", "绩效与考核智能体", "AI 知识库"],
}


def text_of(path):
    from docx import Document
    d = Document(path)
    t = "\n".join(p.text for p in d.paragraphs)
    for tb in d.tables:
        for r in tb.rows:
            t += " | " + " | ".join(c.text for c in r.cells)
    return t, len(d.tables)


def main(folder):
    bad = 0
    for name in EXCLUSIVE:
        path = os.path.join(folder, f"2026-10绩效考核表_{name}.docx")
        if not os.path.exists(path):
            print(f"[缺文件] {name} -> {path}")
            bad += 1
            continue
        t, ntab = text_of(path)
        others = [k for other, keys in EXCLUSIVE.items() if other != name
                  for k in keys if k in t]
        problems = []
        if t.count("**"):
            problems.append(f"残留 ** x{t.count('**')}")
        if t.count("`"):
            problems.append(f"残留反引号 x{t.count('`')}")
        if ntab < 4:
            problems.append(f"表格仅 {ntab} 张")
        if name not in t:
            problems.append("正文缺本人姓名")
        if others:
            problems.append(f"串人: {others}")
        print(("[OK]  " if not problems else "[FAIL]") + f" {name} 表{ntab}张 " +
              ("; ".join(problems) if problems else ""))
        bad += 1 if problems else 0
    print("\n结论：", "全部合格" if bad == 0 else f"{bad} 份不合格，必须重生成")
    return 0 if bad == 0 else 1


if __name__ == "__main__":
    folder = sys.argv[1] if len(sys.argv) > 1 else r"D:\Documents\我的文档\02_团队\考核与绩效\月度成绩\2026年10月\考核下发_2026-10"
    sys.exit(main(folder))
