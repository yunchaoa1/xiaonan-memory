# -*- coding: utf-8 -*-
"""Markdown -> Word 转换（给领导的汇报用）

用法:
    python md2docx_report.py <源.md> <目标.docx>

依赖: python-docx（本机已装 1.2.0；不需要 pandoc）

支持:
  # / ## / ### / ####  标题层级
  | a | b |            表格（自动识别分隔行 |---|，首行作表头加粗底纹）
  **粗体**              行内加粗
  > 引用                引用块（缩进+斜体）
  - / 1. 列表           项目符号 / 编号
  ---                  分隔线（转成一条细横线段落）
中文统一用微软雅黑（同时设 w:eastAsia，否则 Word 里会回退成宋体）。
"""
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Microsoft YaHei"


def set_run_font(run, bold=False, size=None, color=None, italic=False):
    run.font.name = FONT
    # 中文字体必须单独设 eastAsia，否则 Word 回退宋体
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if bold:
        run.bold = True
    if italic:
        run.italic = True
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color


def add_inline(paragraph, text, base_bold=False):
    """解析 **粗体**，其余按普通 run 加进去。"""
    for part in re.split(r"(\*\*.+?\*\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            r = paragraph.add_run(part[2:-2])
            set_run_font(r, bold=True)
        else:
            r = paragraph.add_run(part)
            set_run_font(r, bold=base_bold)


def parse_row(line):
    cells = line.strip().strip("|").split("|")
    return [c.strip() for c in cells]


def is_separator(line):
    return bool(re.match(r"^\s*\|[\s:|-]+\|\s*$", line)) and "-" in line


def build_table(doc, rows):
    if not rows:
        return
    header, body = rows[0], rows[1:]
    ncol = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=ncol)
    table.style = "Table Grid"

    hr = table.add_row().cells
    for i in range(ncol):
        cell = hr[i]
        cell.text = ""
        p = cell.paragraphs[0]
        add_inline(p, header[i] if i < len(header) else "", base_bold=True)
        # 表头浅灰底纹
        shd = cell._tc.get_or_add_tcPr().makeelement(
            qn("w:shd"), {qn("w:val"): "clear", qn("w:fill"): "EDEDED"}
        )
        cell._tc.get_or_add_tcPr().append(shd)

    for row in body:
        cells = table.add_row().cells
        for i in range(ncol):
            cell = cells[i]
            cell.text = ""
            add_inline(cell.paragraphs[0], row[i] if i < len(row) else "")

    doc.add_paragraph()


def convert(src, dst):
    with open(src, encoding="utf-8") as f:
        lines = f.read().splitlines()

    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(2.4)
        s.bottom_margin = Cm(2.4)
        s.left_margin = Cm(2.6)
        s.right_margin = Cm(2.6)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.35

    for lvl in range(1, 5):
        st = doc.styles[f"Heading {lvl}"]
        st.font.name = FONT
        st.font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # 表格：连续的 | 开头行
        if stripped.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not is_separator(lines[i]):
                    block.append(parse_row(lines[i]))
                i += 1
            build_table(doc, block)
            continue

        # 分隔线
        if re.match(r"^-{3,}$", stripped):
            p = doc.add_paragraph()
            p.prs = None if False else p._p.get_or_add_pPr()
            pPr = p._p.get_or_add_pPr()
            pbdr = pPr.makeelement(qn("w:pBdr"), {})
            bottom = pPr.makeelement(
                qn("w:bottom"),
                {qn("w:val"): "single", qn("w:sz"): "6", qn("w:space"): "1", qn("w:color"): "CCCCCC"},
            )
            pbdr.append(bottom)
            pPr.append(pbdr)
            i += 1
            continue

        # 标题
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            lvl = len(m.group(1))
            h = doc.add_heading(level=lvl)
            r = h.add_run(m.group(2).strip())
            size = {1: 18, 2: 14.5, 3: 12.5, 4: 11.5}[lvl]
            set_run_font(r, bold=True, size=size)
            i += 1
            continue

        # 引用
        if stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            add_inline(p, stripped.lstrip("> ").strip())
            for r in p.runs:
                r.italic = True
                r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            i += 1
            continue

        # 列表
        m = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, m.group(2))
            i += 1
            continue
        m = re.match(r"^(\s*)\d+[.)]\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Number")
            add_inline(p, m.group(2))
            i += 1
            continue

        p = doc.add_paragraph()
        add_inline(p, stripped)
        i += 1

    doc.save(dst)
    print(f"saved: {dst}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    convert(sys.argv[1], sys.argv[2])
